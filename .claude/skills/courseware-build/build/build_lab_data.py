#!/usr/bin/env python3
"""Write every lab dataset from lab_data.py to labs/<lab-folder>/data/*.xlsx.

Each workbook carries a titled header block (what the data is, which lab owns it,
the course/version) above a frozen, filtered header row, so a learner who opens
the file in Excel knows what they are looking at without the lab sheet open.
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import course_data as C
import lab_data as L
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=BRAND)
TITLE_F = Font(name="Calibri", size=14, bold=True, color=BRAND)
SUB_F = Font(name="Calibri", size=10, color="595959")
HDR_F = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
NOTE_F = Font(name="Calibri", size=9, color="595959", italic=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_workbook(ds, path):
    wb = Workbook()
    ws = wb.active
    ws.title = ds["title"][:31]

    ncol = len(ds["headers"])
    last = get_column_letter(ncol)

    # ---- header block
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = ds["title"]
    ws["A1"].font = TITLE_F

    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = ds["desc"]
    ws["A2"].font = SUB_F
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    ws.merge_cells(f"A3:{last}3")
    ws["A3"] = (f"Lab {ds['lab']} · {C.TITLE} · {C.COURSE_CODE} · "
                f"Version {C.VERSION}")
    ws["A3"].font = SUB_F

    # ---- table
    HEAD_ROW = 5
    for j, h in enumerate(ds["headers"], 1):
        c = ws.cell(row=HEAD_ROW, column=j, value=h)
        c.font = HDR_F
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[HEAD_ROW].height = 30

    for i, row in enumerate(ds["rows"], HEAD_ROW + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDER
            if isinstance(v, (int, float)):
                c.alignment = Alignment(horizontal="right")

    # ---- column widths from content
    for j, h in enumerate(ds["headers"], 1):
        vals = [str(h)] + [str(r[j - 1]) for r in ds["rows"][:200] if j <= len(r)]
        w = max(len(v) for v in vals) + 3
        ws.column_dimensions[get_column_letter(j)].width = min(max(w, 10), 52)

    # freeze the header and enable filtering — these files get sorted a lot
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=1)
    ws.auto_filter.ref = f"A{HEAD_ROW}:{last}{HEAD_ROW + len(ds['rows'])}"

    # ---- notes under the table
    if ds["notes"]:
        r = HEAD_ROW + len(ds["rows"]) + 2
        ws.cell(row=r, column=1, value="Notes").font = Font(
            name="Calibri", size=10, bold=True, color=BRAND)
        for k, n in enumerate(ds["notes"], 1):
            ws.merge_cells(start_row=r + k, start_column=1,
                           end_row=r + k, end_column=ncol)
            c = ws.cell(row=r + k, column=1, value=f"• {n}")
            c.font = NOTE_F
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


def main(labs_dir, folder_for):
    written = []
    for ds in L.DATASETS.values():
        d = os.path.join(labs_dir, folder_for[ds["lab"]], "data")
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, L.filename(ds, "xlsx"))
        write_workbook(ds, path)
        written.append(path)
    return written


if __name__ == "__main__":
    import build_labs  # noqa: F401  (build_labs runs the whole pipeline)
